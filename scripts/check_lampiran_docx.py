import sys
import os
import re
import json

try:
    import docx
    from docx.shared import Mm, Pt, Inches, Length
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
except ImportError:
    print(json.dumps({"error": "python-docx is not installed"}))
    sys.exit(1)


def get_len_mm(val):
    if val is None:
        return 0.0
    if hasattr(val, 'mm'):
        return round(float(val.mm), 1)
    # If integer in dxa (1 pt = 20 dxa, 1 mm = 56.6929 dxa)
    return round(float(val) / 56.6929, 1)


def get_len_cm(val):
    if val is None:
        return 0.0
    if hasattr(val, 'cm'):
        return round(float(val.cm), 2)
    return round(float(val) / 566.929, 2)


def check_lampiran_docx(input_path, expected_template=None):
    if not os.path.exists(input_path):
        return {
            "status": "Tidak dapat diperiksa",
            "status_code": "CANNOT_CHECK",
            "error": f"File {input_path} tidak ditemukan."
        }

    try:
        doc = docx.Document(input_path)
    except Exception as e:
        return {
            "status": "Tidak dapat diperiksa",
            "status_code": "CANNOT_CHECK",
            "error": f"Gagal membaca file DOCX: {str(e)}"
        }

    # 1. PAGE SETUP (Paper Size & Margins)
    sections = doc.sections
    sec = sections[0] if sections else None

    page_width_mm = get_len_mm(sec.page_width) if sec else 215.0
    page_height_mm = get_len_mm(sec.page_height) if sec else 330.0

    # Determine Paper Name
    if abs(page_width_mm - 215.0) <= 3.0 and abs(page_height_mm - 330.0) <= 3.0:
        paper_name = "F4 / Folio (215 × 330 mm)"
        is_f4 = True
    elif abs(page_width_mm - 210.0) <= 3.0 and abs(page_height_mm - 297.0) <= 3.0:
        paper_name = f"A4 ({int(page_width_mm)} × {int(page_height_mm)} mm)"
        is_f4 = False
    elif abs(page_width_mm - 215.9) <= 3.0 and abs(page_height_mm - 279.4) <= 3.0:
        paper_name = f"Letter ({round(page_width_mm, 1)} × {round(page_height_mm, 1)} mm)"
        is_f4 = False
    elif abs(page_width_mm - 215.9) <= 3.0 and abs(page_height_mm - 355.6) <= 3.0:
        paper_name = f"Legal ({round(page_width_mm, 1)} × {round(page_height_mm, 1)} mm)"
        is_f4 = False
    else:
        paper_name = f"{round(page_width_mm, 1)} × {round(page_height_mm, 1)} mm"
        is_f4 = False

    top_cm = get_len_cm(sec.top_margin) if sec else 2.0
    bottom_cm = get_len_cm(sec.bottom_margin) if sec else 2.0
    left_cm = get_len_cm(sec.left_margin) if sec else 2.0
    right_cm = get_len_cm(sec.right_margin) if sec else 2.0

    is_margin_top_ok = abs(top_cm - 2.0) <= 0.15
    is_margin_bottom_ok = abs(bottom_cm - 2.0) <= 0.15
    is_margin_left_ok = abs(left_cm - 2.0) <= 0.15
    is_margin_right_ok = abs(right_cm - 2.0) <= 0.15

    is_margins_ok = is_margin_top_ok and is_margin_bottom_ok and is_margin_left_ok and is_margin_right_ok

    if top_cm == bottom_cm == left_cm == right_cm:
        margin_found_str = f"{top_cm} cm"
    else:
        margin_found_str = f"Atas: {top_cm} cm, Bawah: {bottom_cm} cm, Kiri: {left_cm} cm, Kanan: {right_cm} cm"

    # 2. TYPOGRAPHY (Font Family, Font Size, Bold)
    fonts_found = set()
    font_sizes_found = set()
    has_bold = False

    for p in doc.paragraphs:
        p_xml = p._element.xml
        if "<w:b/>" in p_xml or '<w:b w:val="true"/>' in p_xml or '<w:b w:val="1"/>' in p_xml or "font-weight:bold" in p_xml:
            has_bold = True

        for r in p.runs:
            r_font = r.font.name
            if r_font:
                fonts_found.add(r_font)
            
            r_size = r.font.size
            if r_size:
                size_pt = round(r_size.pt, 1)
                font_sizes_found.add(size_pt)

            if r.bold is True or r.font.bold is True:
                has_bold = True

    # Check tables for fonts, sizes, and bolds
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_xml = cell._element.xml
                if "<w:b/>" in cell_xml or '<w:b w:val="true"/>' in cell_xml or '<w:b w:val="1"/>' in cell_xml:
                    has_bold = True

                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.name:
                            fonts_found.add(r.font.name)
                        if r.font.size:
                            font_sizes_found.add(round(r.font.size.pt, 1))
                        if r.bold is True or r.font.bold is True:
                            has_bold = True

    if not fonts_found:
        fonts_found_str = "Bookman Old Style (Default)"
        is_font_ok = True
    else:
        non_bookman = [f for f in fonts_found if "bookman" not in f.lower()]
        is_font_ok = len(non_bookman) == 0
        fonts_found_str = ", ".join(sorted(list(fonts_found)))

    if not font_sizes_found:
        sizes_found_str = "12 pt"
        is_size_ok = True
    else:
        non_12pt = [s for s in font_sizes_found if (s < 10.5 or s > 12.5)]
        is_size_ok = len(non_12pt) == 0
        sizes_found_str = ", ".join(f"{s} pt" for s in sorted(list(font_sizes_found)))

    # 3. STRUCTURE (Cover, Front Matter, Header, Footer, Starts from BAB I)
    has_cover = False
    has_kata_pengantar = False
    has_daftar_isi = False
    has_daftar_tabel = False

    forbidden_cover_kw = [r'^\s*cover\b', r'^\s*cover\s+dokumen', r'^\s*lembar\s+pengesahan\b']
    forbidden_pengantar_kw = [r'^\s*kata\s+pengantar\b', r'^\s*preface\b']
    forbidden_toc_kw = [r'^\s*daftar\s+isi\b', r'^\s*table\s+of\s+contents\b']
    forbidden_lot_kw = [r'^\s*daftar\s+tabel\b', r'^\s*daftar\s+gambar\b', r'^\s*daftar\s+grafik\b']

    paragraphs_text = [p.text.strip().lower() for p in doc.paragraphs if p.text.strip()]

    # Scan early paragraphs for front matter before BAB I
    for txt in paragraphs_text[:20]:
        if re.search(r'^\s*bab\s+i\b', txt) or re.search(r'^\s*1\.1\b', txt) or re.search(r'^\s*pendahuluan\b', txt):
            break

        if any(re.search(kw, txt) for kw in forbidden_cover_kw):
            has_cover = True
        if any(re.search(kw, txt) for kw in forbidden_pengantar_kw):
            has_kata_pengantar = True
        if any(re.search(kw, txt) for kw in forbidden_toc_kw):
            has_daftar_isi = True
        if any(re.search(kw, txt) for kw in forbidden_lot_kw):
            has_daftar_tabel = True

    # Check Headers & Footers
    has_header = False
    has_footer = False

    for s in doc.sections:
        if s.header and not s.header.is_linked_to_previous:
            header_text = "".join([hp.text.strip() for hp in s.header.paragraphs]).strip()
            if header_text != "":
                has_header = True
        if s.footer and not s.footer.is_linked_to_previous:
            footer_text = "".join([fp.text.strip() for fp in s.footer.paragraphs]).strip()
            if footer_text != "":
                has_footer = True

    # Check Starts from BAB I
    starts_with_bab1 = False
    first_meaningful = ""
    header_title_kw = ["lampiran", "peraturan bupati", "keputusan bupati", "nomor", "tentang", "rencana kerja"]

    for txt in paragraphs_text[:15]:
        if any(kw in txt for kw in header_title_kw):
            continue
        first_meaningful = txt
        if re.search(r'^\s*bab\s+i\b', txt) or re.search(r'^\s*1\.1\b', txt) or re.search(r'^\s*pendahuluan\b', txt):
            starts_with_bab1 = True
        break

    if not first_meaningful and paragraphs_text:
        first_meaningful = paragraphs_text[0]
        if re.search(r'^\s*bab\s+i\b', first_meaningful):
            starts_with_bab1 = True

    is_structure_ok = (not has_cover) and (not has_kata_pengantar) and (not has_daftar_isi) and (not has_daftar_tabel) and (not has_header) and (not has_footer) and starts_with_bab1

    # 4. TEMPLATE TYPE DETECTION (RENJA_LAMPIRAN_MURNI vs RENJA_LAMPIRAN_PERUBAHAN)
    full_text = "\n".join(paragraphs_text)
    is_perubahan = ("keputusan bupati" in full_text[:2000] or "kepbup" in full_text[:2000] or "perubahan" in full_text[:2000])
    
    if is_perubahan:
        detected_template = "RENJA_LAMPIRAN_PERUBAHAN"
        template_label = "RENJA Lampiran Perubahan (Kepbup)"
    else:
        detected_template = "RENJA_LAMPIRAN_MURNI"
        template_label = "RENJA Lampiran Murni (Perbup)"

    # 5. TABLE INSPECTION
    table_issues = []
    printable_width_mm = page_width_mm - left_cm * 10 - right_cm * 10  # 175 mm for F4 2cm margin
    table_count = len(doc.tables)
    tables_exceed_margin = 0
    tables_too_wide = 0
    tables_cut_off = 0
    tables_overflow = 0

    for idx, tbl in enumerate(doc.tables, start=1):
        tbl_width_mm = 0.0
        col_widths = []
        for col in tbl.columns:
            try:
                w_mm = get_len_mm(col.width)
                col_widths.append(w_mm)
            except Exception:
                pass

        if col_widths and sum(col_widths) > 0:
            tbl_width_mm = sum(col_widths)
        else:
            row_widths = []
            for r in tbl.rows[:3]:
                r_w = sum([get_len_mm(cell.width) for cell in r.cells])
                row_widths.append(r_w)
            tbl_width_mm = max(row_widths) if row_widths else 0.0

        exceeds_margin = tbl_width_mm > (printable_width_mm + 2.0)
        too_wide = tbl_width_mm > (printable_width_mm + 2.0)
        cut_off = tbl_width_mm > (page_width_mm - left_cm * 10)
        overflow = tbl_width_mm > (printable_width_mm + 5.0)

        if exceeds_margin or too_wide or cut_off or overflow:
            if exceeds_margin:
                tables_exceed_margin += 1
            if too_wide:
                tables_too_wide += 1
            if cut_off:
                tables_cut_off += 1
            if overflow:
                tables_overflow += 1

            table_issues.append({
                "table_index": idx,
                "width_mm": round(tbl_width_mm, 1),
                "printable_width_mm": round(printable_width_mm, 1),
                "exceeds_margin": exceeds_margin,
                "too_wide": too_wide,
                "cut_off": cut_off,
                "overflow": overflow,
                "detail": f"Tabel {idx}: Lebar {round(tbl_width_mm, 1)} mm melebihi batas cetak ({round(printable_width_mm, 1)} mm)"
            })

    has_table_issues = len(table_issues) > 0

    # 6. EVALUATE ALL RULES & DETERMINE FINAL STATUS
    rules_items = [
        {
            "id": "paper_size",
            "component": "Ukuran Kertas",
            "standard": "F4 / Folio (215 × 330 mm)",
            "found": paper_name,
            "is_valid": is_f4,
            "status": "✓ Sesuai" if is_f4 else f"✕ Ditemukan {paper_name}",
            "detail": "Kertas F4 / Folio 215 × 330 mm" if is_f4 else f"Ukuran kertas dokumen saat ini: {paper_name}"
        },
        {
            "id": "margin",
            "component": "Margin Halaman",
            "standard": "2 cm (Atas, Bawah, Kiri, Kanan)",
            "found": margin_found_str,
            "is_valid": is_margins_ok,
            "status": "✓ Sesuai" if is_margins_ok else f"✕ Ditemukan {margin_found_str}",
            "detail": "Margin 2 cm seragam di seluruh sisi" if is_margins_ok else f"Margin tidak standar: {margin_found_str}"
        },
        {
            "id": "font_family",
            "component": "Font / Tipografi",
            "standard": "Bookman Old Style",
            "found": fonts_found_str,
            "is_valid": is_font_ok,
            "status": "✓ Sesuai" if is_font_ok else f"✕ Ditemukan {fonts_found_str}",
            "detail": "Jenis font Bookman Old Style" if is_font_ok else f"Terdapat jenis font di luar Bookman Old Style ({fonts_found_str})"
        },
        {
            "id": "font_size",
            "component": "Ukuran Font",
            "standard": "12 pt",
            "found": sizes_found_str,
            "is_valid": is_size_ok,
            "status": "✓ Sesuai" if is_size_ok else f"✕ Ditemukan {sizes_found_str}",
            "detail": "Ukuran font 12 pt seragam" if is_size_ok else f"Ukuran font tidak standar: {sizes_found_str}"
        },
        {
            "id": "bold",
            "component": "Huruf Bold",
            "standard": "Tidak menggunakan Bold",
            "found": "Ditemukan" if has_bold else "Tidak ditemukan",
            "is_valid": not has_bold,
            "status": "✓ Sesuai" if not has_bold else "✕ Ditemukan Bold",
            "detail": "Bebas dari penanda tebal (Bold)" if not has_bold else "Ditemukan penanda/style Bold pada paragraf atau tabel"
        },
        {
            "id": "header",
            "component": "Header",
            "standard": "Tidak ada Header",
            "found": "Ditemukan" if has_header else "Tidak ditemukan",
            "is_valid": not has_header,
            "status": "✓ Sesuai" if not has_header else "✕ Ditemukan Header",
            "detail": "Tanpa Header bawaan" if not has_header else "Ditemukan elemen Header pada dokumen"
        },
        {
            "id": "footer",
            "component": "Footer",
            "standard": "Tidak ada Footer",
            "found": "Ditemukan" if has_footer else "Tidak ditemukan",
            "is_valid": not has_footer,
            "status": "✓ Sesuai" if not has_footer else "✕ Ditemukan Footer",
            "detail": "Tanpa Footer bawaan" if not has_footer else "Ditemukan elemen Footer pada dokumen"
        },
        {
            "id": "structure",
            "component": "Struktur Dokumen",
            "standard": "Dimulai BAB I (Tanpa Cover & TOC)",
            "found": "Dimulai BAB I" if is_structure_ok else "Tidak sesuai (Terdapat Cover/TOC/Front Matter)",
            "is_valid": is_structure_ok,
            "status": "✓ Sesuai" if is_structure_ok else "✕ Struktur Tidak Sesuai",
            "detail": "Dokumen dimulai dari BAB I tanpa Front Matter" if is_structure_ok else "Ditemukan Cover, Kata Pengantar, Daftar Isi/Tabel, atau tidak diawali BAB I"
        },
        {
            "id": "tables",
            "component": "Pemeriksaan Tabel",
            "standard": "Tabel tidak keluar margin / overflow",
            "found": f"{len(table_issues)} dari {table_count} tabel bermasalah" if has_table_issues else f"Seluruh {table_count} tabel presisi",
            "is_valid": not has_table_issues,
            "status": "✓ Sesuai" if not has_table_issues else f"⚠️ Ditemukan {len(table_issues)} Tabel Melebihi Margin",
            "detail": "Seluruh tabel pas dalam margin 17.5 cm" if not has_table_issues else f"Ditemukan {tables_too_wide} tabel terlalu lebar / keluar margin"
        }
    ]

    all_valid = (is_f4 and is_margins_ok and is_font_ok and is_size_ok and (not has_bold) and (not has_header) and (not has_footer) and is_structure_ok and (not has_table_issues))

    if all_valid:
        overall_status = "Sesuai"
        status_code = "COMPLIANT"
    elif has_table_issues and (is_f4 and is_margins_ok and is_font_ok and is_size_ok and (not has_bold) and is_structure_ok):
        overall_status = "Peringatan"
        status_code = "WARNING"
    else:
        overall_status = "Tidak Sesuai"
        status_code = "NON_COMPLIANT"

    report = {
        "status": overall_status,
        "status_code": status_code,
        "template_detected": detected_template,
        "template_label": template_label,
        "is_all_valid": all_valid,
        "summary": {
            "is_f4": is_f4,
            "paper_name": paper_name,
            "margins_ok": is_margins_ok,
            "margin_details": margin_found_str,
            "font_ok": is_font_ok,
            "fonts_found": list(fonts_found),
            "size_ok": is_size_ok,
            "sizes_found": list(font_sizes_found),
            "has_bold": has_bold,
            "has_header": has_header,
            "has_footer": has_footer,
            "structure_ok": is_structure_ok,
            "has_cover": has_cover,
            "has_kata_pengantar": has_kata_pengantar,
            "has_daftar_isi": has_daftar_isi,
            "has_daftar_tabel": has_daftar_tabel,
            "starts_with_bab1": starts_with_bab1,
            "table_count": table_count,
            "tables_exceed_margin": tables_exceed_margin,
            "tables_too_wide": tables_too_wide,
            "tables_cut_off": tables_cut_off,
            "tables_overflow": tables_overflow,
        },
        "items": rules_items,
        "table_issues": table_issues
    }

    return report


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"error": "Usage: python check_lampiran_docx.py <path_to_docx>"}))
        sys.exit(1)

    filepath = sys.argv[1]
    res = check_lampiran_docx(filepath)
    print(json.dumps(res, indent=2))
