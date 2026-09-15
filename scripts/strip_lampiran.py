import sys
import os
import json
import re
import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENTATION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def strip_and_format_for_lampiran(source_docx_path: str, output_docx_path: str, meta: dict):
    """
    Transformasi dokumen RENJA Murni menjadi format resmi Lampiran Renja:
    a. Kertas F4/Folio (215 x 330 mm)
    b. Margin Halaman: 2 cm (kanan, kiri, atas, bawah) seragam
    c. Font: Bookman Old Style, 12pt (tabel 6.5-9pt proporsional agar presisi di Word/PDF)
    d. Tidak menggunakan huruf Bold di semua halaman
    e. Dokumen dimulai langsung dari BAB I (tanpa Cover, Kata Pengantar, Daftar Isi, Daftar Tabel)
    f. Tanpa Header dan Footer
    g. Sisipkan kepala naskah resmi LAMPIRAN di bagian paling atas sebelum BAB I
    h. Otomatis merapikan format, orientasi halaman tabel, dan perataan teks
    i. Optimalisasi tabel: auto-scaling ke lebar bidang cetak (100% printable area),
       penyesuaian padding, cantSplit, dan tblHeader untuk baris header.
    """
    doc = docx.Document(source_docx_path)
    ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    font_name = 'Bookman Old Style'

    # ── 1. Cari letak BAB I pada dokumen sumber ─────────────────────────────
    bab1_child_idx = None
    for i, c in enumerate(doc.element.body):
        if c.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(c, doc)
            t = p.text.strip().upper()
            if t == 'BAB I' or re.match(r'^BAB\s+I(\s|$)', t):
                bab1_child_idx = i
                break

    if bab1_child_idx is None:
        for i, c in enumerate(doc.element.body):
            if c.tag.endswith('p'):
                p = docx.text.paragraph.Paragraph(c, doc)
                if p.text.strip().upper().startswith('BAB '):
                    bab1_child_idx = i
                    break

    if bab1_child_idx is None:
        bab1_child_idx = 0

    # ── 2. Hapus seluruh elemen sebelum BAB I (Cover, Kata Pengantar, TOC, dll) ─
    body = doc.element.body
    for _ in range(bab1_child_idx):
        body.remove(body[0])

    # ── 3. Sisipkan Kepala Naskah Resmi Lampiran sebelum BAB I ──────────────
    romawi = str(meta.get('nomor_lampiran_romawi', 'LAMPIRAN XXXVIII')).upper().strip()
    if not romawi.startswith('LAMPIRAN'):
        romawi = f"LAMPIRAN {romawi}"

    ta_renja = str(meta.get('tahun_anggaran', '2027'))
    ta_perbup = str(int(ta_renja) - 1) if ta_renja.isdigit() and int(ta_renja) > 2000 else '2026'
    is_perubahan = bool(meta.get('is_perubahan', False))

    header_lines = [
        romawi,
        "PERATURAN BUPATI CIREBON" if not is_perubahan else "KEPUTUSAN BUPATI CIREBON",
        f"NOMOR           TAHUN {ta_perbup}",
        "TENTANG",
        "RENCANA KERJA PERANGKAT" if not is_perubahan else "PERUBAHAN RENCANA KERJA PERANGKAT",
        f"DAERAH TAHUN {ta_renja}"
    ]

    header_paragraphs = []
    for idx, line in enumerate(header_lines):
        p = doc.add_paragraph()
        body.insert(idx, p._p)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Mm(101.6)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(24) if idx == len(header_lines) - 1 else Pt(0)

        r = p.add_run(line)
        r.font.name = font_name
        r.font.size = Pt(12)
        r.bold = False
        rPr = r._r.get_or_add_rPr()
        col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
        rPr.append(col)
        header_paragraphs.append(p)

    # ── 4. Konfigurasi Orientasi Section (Landscape vs Portrait) ─────────────
    # Pada Lampiran Renja:
    # - Tabel 2.1 (Evaluasi Pelaksanaan) dan Matriks Renja Bab IV berorientasi LANDSCAPE
    # - Seluruh bab narasi dan tabel lainnya berorientasi PORTRAIT
    sect_body_indices = []
    for i, c in enumerate(body):
        sect = c.find(f'.//{{{ns}}}sectPr') if c.tag.endswith('p') else (c if c.tag.endswith('sectPr') else None)
        if sect is not None:
            sect_body_indices.append(i)

    landscape_section_indices = set()
    for tbl in doc.tables:
        col_count = len(tbl.columns)
        tbl_text = ''.join(tbl._tbl.itertext()).upper()

        # Deteksi Tabel 2.1 (>= 14 kolom) dan Matriks Renja Bab IV
        is_wide_table = (
            col_count >= 14 or
            'REKAPITULASI EVALUASI' in tbl_text or
            ('SKPD:' in tbl_text and 'URUSAN/BIDANG' in tbl_text)
        )

        if is_wide_table:
            tbl_idx = list(body).index(tbl._tbl)
            for s_i, end_b_idx in enumerate(sect_body_indices):
                if tbl_idx <= end_b_idx:
                    landscape_section_indices.add(s_i)
                    break

    # Pastikan section Matriks Renja Bab IV juga tercover
    if len(doc.tables) >= 9:
        matriks_tbl = doc.tables[8]
        m_idx = list(body).index(matriks_tbl._tbl)
        for s_i, end_b_idx in enumerate(sect_body_indices):
            if m_idx <= end_b_idx:
                landscape_section_indices.add(s_i)
                break

    if not landscape_section_indices:
        landscape_section_indices = {1, 6}

    # ── 5. Atur Ukuran F4/Folio (215 x 330 mm) & Margin 2cm di Semua Halaman ─
    for s_idx, s in enumerate(doc.sections):
        if s_idx in landscape_section_indices:
            s.page_width = Mm(330)
            s.page_height = Mm(215)
            s.orientation = WD_ORIENTATION.LANDSCAPE
        else:
            s.page_width = Mm(215)
            s.page_height = Mm(330)
            s.orientation = WD_ORIENTATION.PORTRAIT

        s.top_margin = Mm(20)
        s.bottom_margin = Mm(20)
        s.left_margin = Mm(20)
        s.right_margin = Mm(20)

        # Hapus Header & Footer di semua halaman
        s.header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False
        for p in s.header.paragraphs:
            for r in p.runs:
                r.text = ''
        for p in s.footer.paragraphs:
            for r in p.runs:
                r.text = ''

    # ── 6. Hapus BOLD di Seluruh Document (Styles & Numbering) ───────────────
    try:
        styles_xml = doc.styles.element
        for b in list(styles_xml.iter(f'{{{ns}}}b')):
            b.getparent().remove(b)
        for bCs in list(styles_xml.iter(f'{{{ns}}}bCs')):
            bCs.getparent().remove(bCs)
    except Exception:
        pass

    try:
        if hasattr(doc.part, 'numbering_part') and doc.part.numbering_part is not None:
            num_xml = doc.part.numbering_part.element
            for b in list(num_xml.iter(f'{{{ns}}}b')):
                b.getparent().remove(b)
            for bCs in list(num_xml.iter(f'{{{ns}}}bCs')):
                bCs.getparent().remove(bCs)
    except Exception:
        pass

    # ── 6b. Bersihkan Signature / Tanda Tangan Lama di Akhir Dokumen ───────
    while len(doc.paragraphs) > 0:
        last_p = doc.paragraphs[-1]
        txt = last_p.text.strip().upper()
        if not txt:
            p_elem = last_p._p
            p_elem.getparent().remove(p_elem)
            continue

        is_old_sig = (
            'BUPATI CIREBON' in txt or
            'SEKRETARIS DAERAH' in txt or
            'DIUNDANGKAN DI SUMBER' in txt or
            'PADA TANGGAL' in txt or
            'BERITA DAERAH' in txt or
            'BERITA  DAERAH' in txt or
            'IMRON' in txt or
            'HENDRA NIRMALA' in txt
        )

        if is_old_sig:
            p_elem = last_p._p
            p_elem.getparent().remove(p_elem)
        else:
            break

    # ── 7. Formatting Paragraf (No Bold, Bookman 12pt, Hitam, Perapihan) ────
    for p in doc.paragraphs:
        txt = p.text.strip()
        txt_upper = txt.upper()

        # Ganti heading styles ke Normal
        if p.style.name and any(p.style.name.lower().startswith(x) for x in ['heading', 'judul', 'title']):
            p.style = doc.styles['Normal']

        pPr = p._p.get_or_add_pPr()
        pStyle = pPr.find(f'{{{ns}}}pStyle')
        if pStyle is not None:
            val = pStyle.get(qn('w:val'), '')
            if any(val.lower().startswith(x) for x in ['heading', 'judul', 'title']) or re.match(r'^\d+$', val):
                pStyle.set(qn('w:val'), 'Normal')

        # Hapus shading/background pada paragraf
        for tag in ['shd', 'pBdr']:
            el = pPr.find(f'{{{ns}}}{tag}')
            if el is not None:
                pPr.remove(el)

        # Hapus bold dan warna pada paragraph mark
        pPr_rPr = pPr.find(f'{{{ns}}}rPr')
        if pPr_rPr is not None:
            for tag in ['b', 'bCs', 'color', 'highlight', 'shd']:
                el = pPr_rPr.find(f'{{{ns}}}{tag}')
                if el is not None:
                    pPr_rPr.remove(el)
            b = OxmlElement('w:b'); b.set(qn('w:val'), '0')
            bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '0')
            col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
            pPr_rPr.extend([b, bCs, col])

        # Format setiap run: font Bookman Old Style 12pt, NO BOLD, hitam
        for r in p.runs:
            r.font.name = font_name
            r.font.size = Pt(12)
            r.bold = False
            rPr = r._r.get_or_add_rPr()
            for tag in ['b', 'bCs', 'color', 'highlight', 'shd', 'rStyle']:
                el = rPr.find(f'{{{ns}}}{tag}')
                if el is not None:
                    rPr.remove(el)
            b = OxmlElement('w:b'); b.set(qn('w:val'), '0')
            bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '0')
            col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
            rPr.extend([b, bCs, col])

        # Perapihan alignment otomatis
        if re.match(r'^BAB\s+[IVXLCDM0-9]+$', txt_upper) or txt_upper in [
            'PENDAHULUAN', 'HASIL EVALUASI RENJA', 'TUJUAN DAN SASARAN',
            'RENCANA KERJA DAN PENDANAAN', 'KINERJA PENYELENGGARAAN URUSAN', 'PENUTUP'
        ]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        elif re.match(r'^\d+\.\d+\.?\s+', txt):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif p not in header_paragraphs:
            if p.alignment == WD_ALIGN_PARAGRAPH.LEFT or p.alignment is None:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Anti-orphan BAB Penutup:
    # Set keep_with_next = True pada paragraf terakhir isi BAB Penutup
    # agar tidak terpisah sendirian dari blok tanda tangan
    if len(doc.paragraphs) > 0:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True

    # ── 7b. Injeksi Format Resmi Halaman Penutup (BUPATI CIREBON & IMRON) ────
    p_bupati = doc.add_paragraph()
    p_bupati.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bupati.paragraph_format.left_indent = Mm(101.6)
    p_bupati.paragraph_format.space_before = Pt(24)
    p_bupati.paragraph_format.space_after = Pt(72) # Ruang tanda tangan ~4 baris
    p_bupati.paragraph_format.line_spacing = 1.0
    p_bupati.paragraph_format.keep_with_next = True

    r_bupati = p_bupati.add_run("BUPATI CIREBON,")
    r_bupati.font.name = font_name
    r_bupati.font.size = Pt(12)
    r_bupati.bold = False
    rPr_b = r_bupati._r.get_or_add_rPr()
    for tag in ['b', 'bCs', 'color', 'highlight', 'shd']:
        el = rPr_b.find(f'{{{ns}}}{tag}')
        if el is not None:
            rPr_b.remove(el)
    b_b = OxmlElement('w:b'); b_b.set(qn('w:val'), '0')
    bCs_b = OxmlElement('w:bCs'); bCs_b.set(qn('w:val'), '0')
    col_b = OxmlElement('w:color'); col_b.set(qn('w:val'), '000000')
    rPr_b.extend([b_b, bCs_b, col_b])

    p_imron = doc.add_paragraph()
    p_imron.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_imron.paragraph_format.left_indent = Mm(101.6)
    p_imron.paragraph_format.space_before = Pt(0)
    p_imron.paragraph_format.space_after = Pt(0)
    p_imron.paragraph_format.line_spacing = 1.0
    p_imron.paragraph_format.keep_with_next = False

    r_imron = p_imron.add_run("IMRON")
    r_imron.font.name = font_name
    r_imron.font.size = Pt(12)
    r_imron.bold = False
    rPr_i = r_imron._r.get_or_add_rPr()
    for tag in ['b', 'bCs', 'color', 'highlight', 'shd']:
        el = rPr_i.find(f'{{{ns}}}{tag}')
        if el is not None:
            rPr_i.remove(el)
    b_i = OxmlElement('w:b'); b_i.set(qn('w:val'), '0')
    bCs_i = OxmlElement('w:bCs'); bCs_i.set(qn('w:val'), '0')
    col_i = OxmlElement('w:color'); col_i.set(qn('w:val'), '000000')
    rPr_i.extend([b_i, bCs_i, col_i])

    # ── 8. Formatting & Scaling Semua Tabel (Fit to Printable Width) ────────
    for tbl in doc.tables:
        tbl_idx = list(body).index(tbl._tbl)
        tbl_sect_idx = 0
        for s_i, end_b_idx in enumerate(sect_body_indices):
            if tbl_idx <= end_b_idx:
                tbl_sect_idx = s_i
                break

        is_landscape_section = (tbl_sect_idx in landscape_section_indices)

        # Lebar bidang cetak (printable width) = Lebar Kertas - Margin Kiri - Margin Kanan
        # F4 Landscape (330mm - 40mm = 290mm = 16441 dxa)
        # F4 Portrait (215mm - 40mm = 175mm = 9921 dxa)
        avail_dxa = 16441.0 if is_landscape_section else 9921.0
        num_cols = len(tbl.columns)

        tblPr = tbl._tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement('w:tblPr')
            tbl._tbl.insert(0, tblPr)

        # 8a. Hapus floating positioning agar tabel tidak keluar margin
        tblpPr = tblPr.find(f'{{{ns}}}tblpPr')
        if tblpPr is not None:
            tblPr.remove(tblpPr)

        # 8b. Alignment Left (rata kiri margin 2cm)
        jc = tblPr.find(f'{{{ns}}}jc')
        if jc is None:
            jc = OxmlElement('w:jc')
            tblPr.append(jc)
        jc.set(qn('w:val'), 'left')

        # 8c. tblInd = 0 (tanpa indent negatif)
        tblInd = tblPr.find(f'{{{ns}}}tblInd')
        if tblInd is None:
            tblInd = OxmlElement('w:tblInd')
            tblPr.append(tblInd)
        tblInd.set(qn('w:w'), '0')
        tblInd.set(qn('w:type'), 'dxa')

        # 8d. tblW = avail_dxa
        tblW = tblPr.find(f'{{{ns}}}tblW')
        if tblW is None:
            tblW = OxmlElement('w:tblW')
            tblPr.append(tblW)
        tblW.set(qn('w:w'), str(int(avail_dxa)))
        tblW.set(qn('w:type'), 'dxa')

        # 8e. tblCellMar (padding tabel bersih dan rapi)
        tblCellMar = tblPr.find(f'{{{ns}}}tblCellMar')
        if tblCellMar is None:
            tblCellMar = OxmlElement('w:tblCellMar')
            tblPr.append(tblCellMar)
        for side, val in [('top', '30'), ('bottom', '30'), ('left', '45'), ('right', '45')]:
            node = tblCellMar.find(f'{{{ns}}}{side}')
            if node is None:
                node = OxmlElement(f'w:{side}')
                tblCellMar.append(node)
            node.set(qn('w:w'), val)
            node.set(qn('w:type'), 'dxa')

        # 8f. Proporsi lebar kolom (gridCol & tcW)
        curr_widths = []
        if tbl._tbl.tblGrid is not None and len(tbl._tbl.tblGrid.gridCol_lst) > 0:
            for col in tbl._tbl.tblGrid.gridCol_lst:
                w_val = col.w
                if w_val > 50000:
                    w_val = int(w_val / 56.7)
                curr_widths.append(float(w_val))
        else:
            max_cells_row = max(tbl.rows, key=lambda r: len(r.cells))
            for cell in max_cells_row.cells:
                tcPr = cell._tc.tcPr
                tcW = tcPr.find(f'{{{ns}}}tcW') if tcPr is not None else None
                w_val = float(tcW.get(qn('w:w'), '1000')) if tcW is not None else 1000.0
                curr_widths.append(w_val)

        if not curr_widths or len(curr_widths) != num_cols:
            curr_widths = [avail_dxa / max(1, num_cols)] * num_cols

        tot_w = sum(curr_widths)
        if tot_w <= 0:
            tot_w = avail_dxa
        scale_fact = avail_dxa / tot_w
        new_w_dxa = [int(round(w * scale_fact)) for w in curr_widths]
        diff = int(avail_dxa) - sum(new_w_dxa)
        if new_w_dxa:
            new_w_dxa[-1] += diff

        # Update tblGrid
        if tbl._tbl.tblGrid is not None:
            grid = tbl._tbl.tblGrid
            for col_el in list(grid.gridCol_lst):
                grid.remove(col_el)
            for w_dxa in new_w_dxa:
                col_el = OxmlElement('w:gridCol')
                col_el.set(qn('w:w'), str(w_dxa))
                grid.append(col_el)

        # 8g. Ukuran Font Tabel Proporsional
        if num_cols >= 14:
            table_font_sz = 6.5  # Tabel 2.1 (19-20 kolom)
        elif num_cols >= 8:
            table_font_sz = 8.0  # Matriks Bab IV (8-9 kolom)
        elif num_cols >= 5:
            table_font_sz = 8.5
        else:
            table_font_sz = 9.0

        # 8h. Deteksi baris header
        header_rows_count = 0
        for r_idx, row in enumerate(tbl.rows[:5]):
            row_text = ' '.join(c.text.strip() for c in row.cells)
            if any(kw in row_text.upper() for kw in ['KODE', 'URUSAN', 'INDIKATOR', 'TARGET', 'REALISASI', 'PROGRAM/KEGIATAN', 'ASPEK', 'NO']):
                header_rows_count = r_idx + 1
            elif re.search(r'\b1\b.*\b2\b.*\b3\b', row_text):
                header_rows_count = r_idx + 1
                break

        for r_idx, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            cantSplit = trPr.find(f'{{{ns}}}cantSplit')
            if cantSplit is None:
                cantSplit = OxmlElement('w:cantSplit')
                trPr.append(cantSplit)

            if r_idx < header_rows_count:
                tblHeader = trPr.find(f'{{{ns}}}tblHeader')
                if tblHeader is None:
                    tblHeader = OxmlElement('w:tblHeader')
                    trPr.append(tblHeader)

            c_idx = 0
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                gridSpan = tcPr.find(f'{{{ns}}}gridSpan')
                span = int(gridSpan.get(qn('w:val'), '1')) if gridSpan is not None else 1

                if c_idx < len(new_w_dxa):
                    cell_w_dxa = sum(new_w_dxa[c_idx:c_idx + span])
                    tcW = tcPr.find(f'{{{ns}}}tcW')
                    if tcW is None:
                        tcW = OxmlElement('w:tcW')
                        tcPr.append(tcW)
                    tcW.set(qn('w:w'), str(int(cell_w_dxa)))
                    tcW.set(qn('w:type'), 'dxa')
                c_idx += span

                for p in cell.paragraphs:
                    pPr = p._p.get_or_add_pPr()
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)

                    if r_idx < header_rows_count:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    for r in p.runs:
                        r.font.name = font_name
                        r.font.size = Pt(table_font_sz)
                        r.bold = False

                        rPr = r._r.get_or_add_rPr()
                        for tag in ['b', 'bCs', 'color', 'highlight', 'shd', 'rStyle']:
                            el = rPr.find(f'{{{ns}}}{tag}')
                            if el is not None:
                                rPr.remove(el)
                        b = OxmlElement('w:b'); b.set(qn('w:val'), '0')
                        bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '0')
                        col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
                        rPr.extend([b, bCs, col])

    # ── 9. Simpan Dokumen Hasil Lampiran ─────────────────────────────────────
    os.makedirs(os.path.dirname(output_docx_path), exist_ok=True)
    doc.save(output_docx_path)
    print(f"Successfully generated clean Lampiran DOCX at {output_docx_path}")


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python strip_lampiran.py <meta_json_path> <output_docx_path>")
        print("  meta_json = { source_docx_path, nomor_lampiran_romawi, tahun_anggaran, is_perubahan }")
        sys.exit(1)

    meta_path = sys.argv[1]
    output_path = sys.argv[2]

    with open(meta_path, 'r', encoding='utf-8-sig') as f:
        meta_data = json.load(f)

    src_path = meta_data.get('source_docx_path', '')
    if not src_path or not os.path.exists(src_path):
        print(f"Error: source_docx_path tidak ditemukan: {src_path}", file=sys.stderr)
        sys.exit(1)

    strip_and_format_for_lampiran(src_path, output_path, meta_data)
