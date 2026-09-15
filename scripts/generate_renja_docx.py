import sys
import os
import json
import re
from html.parser import HTMLParser

try:
    import docx
    from docx.shared import Mm, Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    print("Error: python-docx is not installed.", file=sys.stderr)
    sys.exit(1)


class HTMLToDocxParser(HTMLParser):
    def __init__(self, doc, default_font="Bookman Old Style", default_size_pt=12, allow_bold=True):
        super().__init__()
        self.doc = doc
        self.default_font = default_font
        self.default_size_pt = default_size_pt
        self.allow_bold = allow_bold
        self.current_p = None
        self.bold = False
        self.italic = False
        self.underline = False
        self.in_table = False
        self.table = None
        self.current_row = None
        self.current_cell = None
        self.cell_paragraphs = []
        self.is_th = False

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag in ['p', 'div']:
            align_val = WD_ALIGN_PARAGRAPH.JUSTIFY
            attrs_dict = dict(attrs)
            if 'style' in attrs_dict:
                s = attrs_dict['style'].lower()
                if 'center' in s:
                    align_val = WD_ALIGN_PARAGRAPH.CENTER
                elif 'right' in s:
                    align_val = WD_ALIGN_PARAGRAPH.RIGHT
                elif 'left' in s:
                    align_val = WD_ALIGN_PARAGRAPH.LEFT
            
            if self.in_table and self.current_cell:
                self.current_p = self.current_cell.add_paragraph()
            else:
                self.current_p = self.doc.add_paragraph()
            
            self.current_p.alignment = align_val
            self.current_p.paragraph_format.line_spacing = 1.5
            self.current_p.paragraph_format.space_after = Pt(6)
            self.current_p.paragraph_format.space_before = Pt(0)

            # Apply style properties
            if 'style' in attrs_dict:
                s = attrs_dict['style'].lower()
                margin_left_match = re.search(r'margin-left\s*:\s*([\d.]+)(cm|mm|in|pt)', s)
                if margin_left_match:
                    val = float(margin_left_match.group(1))
                    unit = margin_left_match.group(2)
                    if unit == 'cm':
                        self.current_p.paragraph_format.left_indent = Mm(val * 10)
                    elif unit == 'mm':
                        self.current_p.paragraph_format.left_indent = Mm(val)
                    elif unit == 'in':
                        self.current_p.paragraph_format.left_indent = Inches(val)
                    elif unit == 'pt':
                        self.current_p.paragraph_format.left_indent = Pt(val)
                
                if 'keep-together' in s or 'page-break-inside' in s or 'avoid' in s:
                    self.current_p.paragraph_format.keep_with_next = True
            
        elif tag == 'b' or tag == 'strong':
            self.bold = self.allow_bold
        elif tag == 'i' or tag == 'em':
            self.italic = True
        elif tag == 'u':
            self.underline = True
        elif tag == 'br':
            if self.current_p:
                self.current_p.add_run().add_break()
            else:
                p = self.doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.5
        elif tag == 'table':
            self.in_table = True
            # Create a table dynamically
            self.table_rows_data = []
        elif tag == 'tr':
            self.current_row_data = []
        elif tag in ['td', 'th']:
            self.is_th = (tag == 'th')
            self.cell_text = ""

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ['p', 'div']:
            self.current_p = None
        elif tag in ['b', 'strong']:
            self.bold = False
        elif tag in ['i', 'em']:
            self.italic = False
        elif tag == 'u':
            self.underline = False
        elif tag in ['td', 'th']:
            if hasattr(self, 'current_row_data'):
                self.current_row_data.append({
                    'text': self.cell_text.strip(),
                    'is_th': self.is_th
                })
        elif tag == 'tr':
            if hasattr(self, 'table_rows_data') and hasattr(self, 'current_row_data'):
                self.table_rows_data.append(self.current_row_data)
        elif tag == 'table':
            self.in_table = False
            self.render_parsed_table()

    def handle_data(self, data):
        text = data.replace('\xa0', ' ')
        if not text:
            return

        if self.in_table:
            if hasattr(self, 'cell_text'):
                self.cell_text += text
            return

        if self.current_p is None:
            self.current_p = self.doc.add_paragraph()
            self.current_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self.current_p.paragraph_format.line_spacing = 1.5
            self.current_p.paragraph_format.space_after = Pt(6)

        run = self.current_p.add_run(text)
        run.font.name = self.default_font
        run.font.size = Pt(self.default_size_pt)
        run.bold = self.bold
        run.italic = self.italic
        run.underline = self.underline

    def render_parsed_table(self):
        if not hasattr(self, 'table_rows_data') or not self.table_rows_data:
            return

        num_rows = len(self.table_rows_data)
        num_cols = max(len(row) for row in self.table_rows_data) if num_rows > 0 else 0
        if num_rows == 0 or num_cols == 0:
            return

        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for r_idx, row_data in enumerate(self.table_rows_data):
            for c_idx, cell_data in enumerate(row_data):
                if c_idx < num_cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_data['text']
                    
                    # Set cell styling
                    p = cell.paragraphs[0]
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    
                    if p.runs:
                        for run in p.runs:
                            run.font.name = self.default_font
                            run.font.size = Pt(self.default_size_pt - 1)
                            if cell_data['is_th']:
                                run.bold = self.allow_bold
                    
                    # Background shading for header
                    if cell_data['is_th']:
                        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E2E8F0"/>')
                        cell._tc.get_or_add_tcPr().append(shd)

        # Set borders for table
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(tblBorders)

        # Add spacing after table
        p_after = self.doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)


def build_renja_docx(data, output_path):
    doc = docx.Document()
    doc_type = str(data.get('document_type') or '')
    title = str(data.get('title') or '')
    is_lampiran = bool('lampiran' in doc_type.lower() or 'lampiran' in title.lower() or doc_type in ['RENJA_LAMPIRAN_MURNI', 'RENJA_LAMPIRAN_PERUBAHAN'])
    is_blank_template = bool(data.get('is_blank_template') or data.get('is_template'))
    allow_bold = not is_lampiran
    
    # 1. Page Setup F4 Folio (21.5 cm x 33.0 cm)
    section = doc.sections[0]
    section.page_width = Mm(215)
    section.page_height = Mm(330)
    
    # Official Standard Margins (2.0 cm / 20 mm all around)
    margins_config = data.get('margins') or {}
    section.top_margin = Mm(margins_config.get('top_mm', 20))
    section.bottom_margin = Mm(margins_config.get('bottom_mm', 20))
    section.left_margin = Mm(margins_config.get('left_mm', 20))
    section.right_margin = Mm(margins_config.get('right_mm', 20))

    font_family = "Bookman Old Style"
    font_size_pt = 12

    # Set normal style font
    style_normal = doc.styles['Normal']
    style_normal.font.name = font_family
    style_normal.font.size = Pt(font_size_pt)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)

    # 2. Cover if exists
    cover_data = data.get('cover_data') or {}
    has_cover = bool(cover_data and cover_data.get('judul_dokumen')) and not is_lampiran
    
    if has_cover:
        p_cov = doc.add_paragraph()
        p_cov.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cov.paragraph_format.space_before = Pt(48)
        
        r_title = p_cov.add_run(cover_data.get('judul_dokumen', 'DOKUMEN RENJA').upper() + '\n\n')
        r_title.bold = True
        r_title.font.size = Pt(16)
        
        r_opd = p_cov.add_run((data.get('opd_name') or '{{NAMA_OPD}}').upper() + '\n')
        r_opd.bold = True
        r_opd.font.size = Pt(14)

        r_ta = p_cov.add_run(f"TAHUN ANGGARAN {data.get('tahun_anggaran', '{{TAHUN_ANGGARAN}}')}\n\n")
        r_ta.bold = True
        r_ta.font.size = Pt(12)

        # Optional Logo on Cover
        logo_path = data.get('logo_path')
        if logo_path and os.path.exists(logo_path):
            try:
                p_logo = doc.add_paragraph()
                p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_logo.paragraph_format.space_before = Pt(12)
                p_logo.paragraph_format.space_after = Pt(24)
                r_logo = p_logo.add_run()
                r_logo.add_picture(logo_path, width=Inches(1.8))
            except Exception as e:
                # If image loading fails, proceed gracefully without dummy
                pass
        else:
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_before = Pt(36)

        p_pemda = doc.add_paragraph()
        p_pemda.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_pemda.paragraph_format.space_before = Pt(24)
        p_pemda.paragraph_format.space_after = Pt(0)
        r_pemda = p_pemda.add_run("PEMERINTAH KABUPATEN CIREBON\nSUMBER\n")
        r_pemda.bold = True
        r_pemda.font.size = Pt(14)

        doc.add_page_break()

    # 3. Official Perbup Header Block on first page (Only for Lampiran if not already in content)
    sections_list = data.get('sections') or []
    first_sec_content = sections_list[0].get('content') or '' if sections_list else ''
    
    is_perubahan = bool('perubahan' in doc_type.lower() or 'perubahan' in title.lower() or doc_type == 'RENJA_LAMPIRAN_PERUBAHAN')
    should_render_header = is_lampiran and not ("PERATURAN BUPATI CIREBON" in first_sec_content or "KEPUTUSAN BUPATI CIREBON" in first_sec_content)

    if should_render_header:
        romawi = (data.get('nomor_lampiran_romawi') or 'LAMPIRAN I').upper()
        if not romawi.startswith('LAMPIRAN'):
            romawi = f"LAMPIRAN {romawi}"

        ta_renja = str(data.get('tahun_anggaran') or '2027')
        ta_peraturan = ta_renja if is_perubahan else (str(int(ta_renja) - 1) if ta_renja.isdigit() and int(ta_renja) > 2000 else '2026')

        header_lines = [
            romawi,
            "PERATURAN BUPATI CIREBON" if not is_perubahan else "KEPUTUSAN BUPATI CIREBON",
            f"NOMOR           TAHUN {ta_peraturan}",
            "TENTANG",
            "RENCANA KERJA PERANGKAT DAERAH" if not is_perubahan else "PERUBAHAN RENCANA KERJA PERANGKAT DAERAH",
            f"TAHUN {ta_renja}"
        ]

        for idx, line in enumerate(header_lines):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Mm(101.6)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(24) if idx == len(header_lines) - 1 else Pt(0)

            r = p.add_run(line)
            r.font.name = font_family
            r.font.size = Pt(font_size_pt)
            r.bold = False
            rPr = r._r.get_or_add_rPr()
            col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
            rPr.append(col)

    # 4. Sections Rendering
    # If lampiran, filter out front sections
    filtered_sections_list = []
    for s in sections_list:
        b_code = str(s.get('bab_code') or '').upper().strip()
        b_title = str(s.get('bab_title') or '').lower().strip()
        sec_type = str(s.get('section_type') or '').lower().strip()

        if is_lampiran:
            if b_code in ['COVER', 'PENGESAHAN', 'PREFACE', 'TOC', 'LOT', 'LOF', 'LOC', 'LOA'] or \
               sec_type in ['cover', 'preface', 'table_of_contents', 'list_of_tables', 'list_of_figures', 'list_of_charts', 'list_of_appendices'] or \
               'kata pengantar' in b_title or 'daftar isi' in b_title or 'lembar pengesahan' in b_title or 'cover' in b_title:
                continue
        elif sec_type == 'cover':
            # Cover already rendered via cover_data
            continue

        filtered_sections_list.append(s)

    # Render Front Sections (for non-lampiran) sequentially
    front_types = ['preface', 'table_of_contents', 'list_of_tables', 'list_of_figures', 'list_of_charts', 'list_of_appendices']
    rendered_front_keys = set()

    for s in filtered_sections_list:
        sec_type = str(s.get('section_type') or '').lower().strip()
        if sec_type in front_types:
            b_title = s.get('bab_title') or s.get('sub_bab_title') or s.get('title') or ''
            
            # Add page break before front section
            if s.get('page_break_before', True):
                if len(doc.paragraphs) > 0:
                    doc.add_page_break()

            p_front = doc.add_paragraph()
            p_front.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_front.paragraph_format.space_before = Pt(24)
            p_front.paragraph_format.space_after = Pt(12)
            r_front = p_front.add_run(b_title.upper())
            r_front.bold = True
            r_front.font.name = font_family
            r_front.font.size = Pt(font_size_pt)

            content_html = s.get('content') or ''
            if content_html:
                parser = HTMLToDocxParser(doc, default_font=font_family, default_size_pt=font_size_pt, allow_bold=allow_bold)
                try:
                    parser.feed(content_html)
                except Exception:
                    plain_p = doc.add_paragraph()
                    plain_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    clean_txt = re.sub(r'<[^>]+>', '', content_html).strip()
                    r_plain = plain_p.add_run(clean_txt)
                    r_plain.font.name = font_family
                    r_plain.font.size = Pt(font_size_pt)

    # Group Chapters and Subchapters
    main_sections = [s for s in filtered_sections_list if str(s.get('section_type') or '').lower().strip() not in front_types]

    grouped_babs = {}
    for s in main_sections:
        b_code = s.get('bab_code') or 'BAB I'
        if b_code not in grouped_babs:
            grouped_babs[b_code] = []
        grouped_babs[b_code].append(s)

    is_first_main_section = True
    for b_code, secs in grouped_babs.items():
        first_sec = secs[0] if secs else {}
        b_title = first_sec.get('bab_title') or ''
        sec_type = str(first_sec.get('section_type') or '').lower().strip()

        # Page break before each BAB
        if first_sec.get('page_break_before', True) or (not is_lampiran and not is_first_main_section):
            if len(doc.paragraphs) > 0:
                doc.add_page_break()

        is_first_main_section = False

        if sec_type == 'appendix' or 'lampiran' in b_code.lower():
            # Appendix Section
            p_app = doc.add_paragraph()
            p_app.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_app.paragraph_format.space_before = Pt(24)
            p_app.paragraph_format.space_after = Pt(6)
            r_app = p_app.add_run(b_code.upper())
            r_app.bold = allow_bold
            r_app.font.name = font_family
            r_app.font.size = Pt(font_size_pt)

            if b_title and b_title.upper() != b_code.upper():
                p_app_title = doc.add_paragraph()
                p_app_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_app_title.paragraph_format.space_before = Pt(0)
                p_app_title.paragraph_format.space_after = Pt(12)
                r_app_title = p_app_title.add_run(b_title.upper())
                r_app_title.bold = allow_bold
                r_app_title.font.name = font_family
                r_app_title.font.size = Pt(font_size_pt)

            for s in secs:
                content_html = s.get('content') or ''
                if content_html:
                    parser = HTMLToDocxParser(doc, default_font=font_family, default_size_pt=font_size_pt, allow_bold=allow_bold)
                    try:
                        parser.feed(content_html)
                    except Exception:
                        clean_txt = re.sub(r'<[^>]+>', '', content_html).strip()
                        plain_p = doc.add_paragraph()
                        plain_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        plain_p.add_run(clean_txt)
            continue

        # Standard BAB Heading
        p_bab = doc.add_paragraph()
        p_bab.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_bab.paragraph_format.space_before = Pt(24)
        p_bab.paragraph_format.space_after = Pt(2)
        r_bab_code = p_bab.add_run(b_code.upper())
        r_bab_code.bold = allow_bold
        r_bab_code.font.name = font_family
        r_bab_code.font.size = Pt(font_size_pt)

        if b_title:
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_title.paragraph_format.space_before = Pt(0)
            p_title.paragraph_format.space_after = Pt(12)
            r_title = p_title.add_run(b_title.upper())
            r_title.bold = allow_bold
            r_title.font.name = font_family
            r_title.font.size = Pt(font_size_pt)

        # Subbabs
        for s in secs:
            sub_code = (s.get('sub_bab_code') or '').strip()
            sub_title = (s.get('sub_bab_title') or '').strip()
            clean_title = re.sub(r'^\d+(\.\d+)*\s*', '', sub_title).strip()
            full_sub_heading = f"{sub_code}. {clean_title}".upper() if sub_code else clean_title.upper()

            # Cegah cetak subbab ganda jika sub_code/sub_title mengulang nama BAB
            is_repeating_bab = (sub_code.upper().startswith("BAB ") or clean_title.upper().startswith("BAB ") or
                                (b_code.upper() in sub_code.upper()) or 
                                (b_title and b_title.upper() in sub_title.upper() and len(clean_title) < len(b_title) + 8))

            if full_sub_heading and not is_repeating_bab:
                p_sub = doc.add_paragraph()
                p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p_sub.paragraph_format.space_before = Pt(10)
                p_sub.paragraph_format.space_after = Pt(4)
                r_sub = p_sub.add_run(full_sub_heading)
                r_sub.bold = allow_bold
                r_sub.font.name = font_family
                r_sub.font.size = Pt(font_size_pt)

            content_html = s.get('content') or ''
            if content_html:
                # Hapus paragraf awal pada content_html jika mengulangi judul BAB
                content_html = re.sub(r'^\s*<p\b[^>]*>\s*(?:<b>|<strong>)?\s*BAB\s+[IVXLCDM]+[\.\:]?\s*[^<]*<\/p>', '', content_html, flags=re.IGNORECASE)
                content_html = content_html.strip()

                if content_html:
                    parser = HTMLToDocxParser(doc, default_font=font_family, default_size_pt=font_size_pt, allow_bold=allow_bold)
                    try:
                        parser.feed(content_html)
                    except Exception as e:
                        plain_p = doc.add_paragraph()
                        plain_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        plain_p.paragraph_format.line_spacing = 1.15
                        plain_p.paragraph_format.space_after = Pt(6)
                        clean_txt = re.sub(r'<[^>]+>', '', content_html).strip()
                        r_plain = plain_p.add_run(clean_txt)
                        r_plain.font.name = font_family
                        r_plain.font.size = Pt(font_size_pt)
                        r_plain.bold = False

    # Append Official Closing Signature for Lampiran
    if is_lampiran:
        if len(doc.paragraphs) > 0:
            doc.paragraphs[-1].paragraph_format.keep_with_next = True

        p_bupati = doc.add_paragraph()
        p_bupati.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_bupati.paragraph_format.left_indent = Mm(101.6)
        p_bupati.paragraph_format.space_before = Pt(24)
        p_bupati.paragraph_format.space_after = Pt(72)
        p_bupati.paragraph_format.line_spacing = 1.0
        p_bupati.paragraph_format.keep_with_next = True

        r_bupati = p_bupati.add_run("BUPATI CIREBON,")
        r_bupati.font.name = font_family
        r_bupati.font.size = Pt(font_size_pt)
        r_bupati.bold = False

        p_imron = doc.add_paragraph()
        p_imron.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_imron.paragraph_format.left_indent = Mm(101.6)
        p_imron.paragraph_format.space_before = Pt(0)
        p_imron.paragraph_format.space_after = Pt(0)
        p_imron.paragraph_format.line_spacing = 1.0
        p_imron.paragraph_format.keep_with_next = False

        r_imron = p_imron.add_run("IMRON")
        r_imron.font.name = font_family
        r_imron.font.size = Pt(font_size_pt)
        r_imron.bold = False

    # Save to destination
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Successfully generated DOCX at {output_path}")


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python generate_renja_docx.py <input_json_path> <output_docx_path>")
        sys.exit(1)

    json_path = sys.argv[1]
    out_path = sys.argv[2]

    with open(json_path, 'r', encoding='utf-8') as f:
        doc_data = json.load(f)

    build_renja_docx(doc_data, out_path)
