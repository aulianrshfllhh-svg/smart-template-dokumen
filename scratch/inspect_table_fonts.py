import docx

src = 'storage/app/private/renja/2027/kecamatan-depok/renja-murni/original/6a859b9a61741_1. Renja Akhir 2027 Kec Depok.docx'
doc = docx.Document(src)

for t_idx in [0, 6, 8]:
    tbl = doc.tables[t_idx]
    print(f"\n=== TABLE {t_idx} ===")
    fonts = set()
    sizes = set()
    aligns = set()
    bolds = set()
    for row in tbl.rows[:10]:
        for cell in row.cells:
            for p in cell.paragraphs:
                aligns.add(p.alignment)
                for r in p.runs:
                    fonts.add(r.font.name)
                    sizes.add(r.font.size.pt if r.font.size else 'None')
                    bolds.add(r.bold)
    print(f"Fonts: {fonts}")
    print(f"Sizes: {sizes}")
    print(f"Alignments: {aligns}")
    print(f"Bolds: {bolds}")
