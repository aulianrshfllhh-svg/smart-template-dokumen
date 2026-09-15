import os
import re
import docx
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENTATION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def test_transform():
    src = 'storage/app/private/renja/2027/kecamatan-depok/renja-murni/original/6a859b9a61741_1. Renja Akhir 2027 Kec Depok.docx'
    out = 'storage/app/temp_docs/test_perfect_table_render.docx'
    pdf_out = 'storage/app/temp_docs/test_perfect_table_render.pdf'
    meta = {'nomor_lampiran_romawi': 'LAMPIRAN XXXVIII', 'tahun_anggaran': '2027', 'is_perubahan': False}

    doc = docx.Document(src)
    ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    font_name = 'Bookman Old Style'

    # 1. BAB I
    bab1_child_idx = None
    for i, c in enumerate(doc.element.body):
        if c.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(c, doc)
            t = p.text.strip().upper()
            if t == 'BAB I' or re.match(r'^BAB\s+I(\s|$)', t):
                bab1_child_idx = i
                break

    if bab1_child_idx is None:
        for i, c in enumerate(doc.element.body):
            if c.tag.endswith('p'):
                p = docx.text.paragraph.Paragraph(c, doc)
                if p.text.strip().upper().startswith('BAB '):
                    bab1_child_idx = i
                    break

    if bab1_child_idx is None:
        bab1_child_idx = 0

    body = doc.element.body
    for _ in range(bab1_child_idx):
        body.remove(body[0])

    # 2. Header Block
    romawi = str(meta.get('nomor_lampiran_romawi', 'LAMPIRAN XXXVIII')).upper().strip()
    if not romawi.startswith('LAMPIRAN'):
        romawi = f"LAMPIRAN {romawi}"
    ta_renja = str(meta.get('tahun_anggaran', '2027'))
    ta_perbup = str(int(ta_renja) - 1) if ta_renja.isdigit() and int(ta_renja) > 2000 else '2026'
    is_perubahan = bool(meta.get('is_perubahan', False))

    header_lines = [
        romawi,
        "PERATURAN BUPATI CIREBON" if not is_perubahan else "KEPUTUSAN BUPATI CIREBON",
        f"NOMOR           TAHUN {ta_perbup}",
        "TENTANG",
        "RENCANA KERJA PERANGKAT" if not is_perubahan else "PERUBAHAN RENCANA KERJA PERANGKAT",
        f"DAERAH TAHUN {ta_renja}"
    ]

    header_paragraphs = []
    for idx, line in enumerate(header_lines):
        p = doc.add_paragraph()
        body.insert(idx, p._p)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Mm(101.6)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(24) if idx == len(header_lines) - 1 else Pt(0)

        r = p.add_run(line)
        r.font.name = font_name
        r.font.size = Pt(12)
        r.bold = False
        rPr = r._r.get_or_add_rPr()
        col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
        rPr.append(col)
        header_paragraphs.append(p)

    # 3. Orientasi Section
    sect_body_indices = []
    for i, c in enumerate(body):
        sect = c.find(f'.//{{{ns}}}sectPr') if c.tag.endswith('p') else (c if c.tag.endswith('sectPr') else None)
        if sect is not None:
            sect_body_indices.append(i)

    landscape_section_indices = set()
    for tbl in doc.tables:
        col_count = len(tbl.columns)
        tbl_text = ''.join(tbl._tbl.itertext()).upper()

        is_wide_table = (
            col_count >= 14 or
            'REKAPITULASI EVALUASI' in tbl_text or
            ('SKPD:' in tbl_text and 'URUSAN/BIDANG' in tbl_text)
        )

        if is_wide_table:
            tbl_idx = list(body).index(tbl._tbl)
            for s_i, end_b_idx in enumerate(sect_body_indices):
                if tbl_idx <= end_b_idx:
                    landscape_section_indices.add(s_i)
                    break

    if len(doc.tables) >= 9:
        matriks_tbl = doc.tables[8]
        m_idx = list(body).index(matriks_tbl._tbl)
        for s_i, end_b_idx in enumerate(sect_body_indices):
            if m_idx <= end_b_idx:
                landscape_section_indices.add(s_i)
                break

    if not landscape_section_indices:
        landscape_section_indices = {1, 6}

    # 4. Set Section Margins & Dimensions
    for s_idx, s in enumerate(doc.sections):
        if s_idx in landscape_section_indices:
            s.page_width = Mm(330)
            s.page_height = Mm(215)
            s.orientation = WD_ORIENTATION.LANDSCAPE
        else:
            s.page_width = Mm(215)
            s.page_height = Mm(330)
            s.orientation = WD_ORIENTATION.PORTRAIT

        s.top_margin = Mm(20)
        s.bottom_margin = Mm(20)
        s.left_margin = Mm(20)
        s.right_margin = Mm(20)

        s.header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False
        for p in s.header.paragraphs:
            for r in p.runs:
                r.text = ''
        for p in s.footer.paragraphs:
            for r in p.runs:
                r.text = ''

    # 5. Remove bold styles
    try:
        styles_xml = doc.styles.element
        for b in list(styles_xml.iter(f'{{{ns}}}b')):
            b.getparent().remove(b)
        for bCs in list(styles_xml.iter(f'{{{ns}}}bCs')):
            bCs.getparent().remove(bCs)
    except Exception:
        pass

    try:
        if hasattr(doc.part, 'numbering_part') and doc.part.numbering_part is not None:
            num_xml = doc.part.numbering_part.element
            for b in list(num_xml.iter(f'{{{ns}}}b')):
                b.getparent().remove(b)
            for bCs in list(num_xml.iter(f'{{{ns}}}bCs')):
                bCs.getparent().remove(bCs)
    except Exception:
        pass

    # 6. Format Paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        txt_upper = txt.upper()

        if p.style.name and any(p.style.name.lower().startswith(x) for x in ['heading', 'judul', 'title']):
            p.style = doc.styles['Normal']

        pPr = p._p.get_or_add_pPr()
        pStyle = pPr.find(f'{{{ns}}}pStyle')
        if pStyle is not None:
            val = pStyle.get(qn('w:val'), '')
            if any(val.lower().startswith(x) for x in ['heading', 'judul', 'title']) or re.match(r'^\d+$', val):
                pStyle.set(qn('w:val'), 'Normal')

        for tag in ['shd', 'pBdr']:
            el = pPr.find(f'{{{ns}}}{tag}')
            if el is not None:
                pPr.remove(el)

        pPr_rPr = pPr.find(f'{{{ns}}}rPr')
        if pPr_rPr is not None:
            for tag in ['b', 'bCs', 'color', 'highlight', 'shd']:
                el = pPr_rPr.find(f'{{{ns}}}{tag}')
                if el is not None:
                    pPr_rPr.remove(el)
            b = OxmlElement('w:b'); b.set(qn('w:val'), '0')
            bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '0')
            col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
            pPr_rPr.extend([b, bCs, col])

        for r in p.runs:
            r.font.name = font_name
            r.font.size = Pt(12)
            r.bold = False
            rPr = r._r.get_or_add_rPr()
            for tag in ['b', 'bCs', 'color', 'highlight', 'shd', 'rStyle']:
                el = rPr.find(f'{{{ns}}}{tag}')
                if el is not None:
                    rPr.remove(el)
            b = OxmlElement('w:b'); b.set(qn('w:val'), '0')
            bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '0')
            col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
            rPr.extend([b, bCs, col])

        if re.match(r'^BAB\s+[IVXLCDM0-9]+$', txt_upper) or txt_upper in [
            'PENDAHULUAN', 'HASIL EVALUASI RENJA', 'TUJUAN DAN SASARAN',
            'RENCANA KERJA DAN PENDANAAN', 'KINERJA PENYELENGGARAAN URUSAN', 'PENUTUP'
        ]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        elif re.match(r'^\d+\.\d+\.?\s+', txt):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif p not in header_paragraphs:
            if p.alignment == WD_ALIGN_PARAGRAPH.LEFT or p.alignment is None:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 7. Format Tables (Optimized for exact page printable width)
    for tbl in doc.tables:
        tbl_idx = list(body).index(tbl._tbl)
        tbl_sect_idx = 0
        for s_i, end_b_idx in enumerate(sect_body_indices):
            if tbl_idx <= end_b_idx:
                tbl_sect_idx = s_i
                break

        is_landscape_section = (tbl_sect_idx in landscape_section_indices)
        
        # Available width:
        # F4 Landscape: 330mm - 40mm = 290mm = 16441 dxa
        # F4 Portrait: 215mm - 40mm = 175mm = 9921 dxa
        avail_dxa = 16441.0 if is_landscape_section else 9921.0
        
        num_cols = len(tbl.columns)
        
        # Setting tblPr
        tblPr = tbl._tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement('w:tblPr')
            tbl._tbl.insert(0, tblPr)

        # Hapus floating positioning
        tblpPr = tblPr.find(f'{{{ns}}}tblpPr')
        if tblpPr is not None:
            tblPr.remove(tblpPr)

        # Alignment left
        jc = tblPr.find(f'{{{ns}}}jc')
        if jc is None:
            jc = OxmlElement('w:jc')
            tblPr.append(jc)
        jc.set(qn('w:val'), 'left')

        # tblInd = 0
        tblInd = tblPr.find(f'{{{ns}}}tblInd')
        if tblInd is None:
            tblInd = OxmlElement('w:tblInd')
            tblPr.append(tblInd)
        tblInd.set(qn('w:w'), '0')
        tblInd.set(qn('w:type'), 'dxa')

        # tblW
        tblW = tblPr.find(f'{{{ns}}}tblW')
        if tblW is None:
            tblW = OxmlElement('w:tblW')
            tblPr.append(tblW)
        tblW.set(qn('w:w'), str(int(avail_dxa)))
        tblW.set(qn('w:type'), 'dxa')

        # tblCellMar
        tblCellMar = tblPr.find(f'{{{ns}}}tblCellMar')
        if tblCellMar is None:
            tblCellMar = OxmlElement('w:tblCellMar')
            tblPr.append(tblCellMar)
        for side, val in [('top', '30'), ('bottom', '30'), ('left', '45'), ('right', '45')]:
            node = tblCellMar.find(f'{{{ns}}}{side}')
            if node is None:
                node = OxmlElement(f'w:{side}')
                tblCellMar.append(node)
            node.set(qn('w:w'), val)
            node.set(qn('w:type'), 'dxa')

        # Ambil lebar kolom asli untuk proporsi
        curr_widths = []
        if tbl._tbl.tblGrid is not None and len(tbl._tbl.tblGrid.gridCol_lst) > 0:
            for col in tbl._tbl.tblGrid.gridCol_lst:
                w_val = col.w
                if w_val > 50000:
                    w_val = int(w_val / 56.7)
                curr_widths.append(float(w_val))
        else:
            max_cells_row = max(tbl.rows, key=lambda r: len(r.cells))
            for cell in max_cells_row.cells:
                tcPr = cell._tc.tcPr
                tcW = tcPr.find(f'{{{ns}}}tcW') if tcPr is not None else None
                w_val = float(tcW.get(qn('w:w'), '1000')) if tcW is not None else 1000.0
                curr_widths.append(w_val)

        if not curr_widths or len(curr_widths) != num_cols:
            curr_widths = [avail_dxa / max(1, num_cols)] * num_cols

        tot_w = sum(curr_widths)
        if tot_w <= 0:
            tot_w = avail_dxa
        scale_fact = avail_dxa / tot_w
        new_w_dxa = [int(round(w * scale_fact)) for w in curr_widths]
        diff = int(avail_dxa) - sum(new_w_dxa)
        if new_w_dxa:
            new_w_dxa[-1] += diff

        # Update tblGrid
        if tbl._tbl.tblGrid is not None:
            grid = tbl._tbl.tblGrid
            for col_el in list(grid.gridCol_lst):
                grid.remove(col_el)
            for w_dxa in new_w_dxa:
                col_el = OxmlElement('w:gridCol')
                col_el.set(qn('w:w'), str(w_dxa))
                grid.append(col_el)

        # Font size for table
        if num_cols >= 14:
            table_font_sz = 6.5 # Sized for 19-20 column wide tables like Table 2.1
        elif num_cols >= 8:
            table_font_sz = 8.0 # Sized for Matriks Renja Bab IV
        elif num_cols >= 5:
            table_font_sz = 8.5
        else:
            table_font_sz = 9.0

        # Deteksi baris header
        header_rows_count = 0
        for r_idx, row in enumerate(tbl.rows[:5]):
            row_text = ' '.join(c.text.strip() for c in row.cells)
            if any(kw in row_text.upper() for kw in ['KODE', 'URUSAN', 'INDIKATOR', 'TARGET', 'REALISASI', 'PROGRAM/KEGIATAN', 'ASPEK', 'NO']):
                header_rows_count = r_idx + 1
            elif re.search(r'\b1\b.*\b2\b.*\b3\b', row_text):
                header_rows_count = r_idx + 1
                break

        for r_idx, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            cantSplit = trPr.find(f'{{{ns}}}cantSplit')
            if cantSplit is None:
                cantSplit = OxmlElement('w:cantSplit')
                trPr.append(cantSplit)

            if r_idx < header_rows_count:
                tblHeader = trPr.find(f'{{{ns}}}tblHeader')
                if tblHeader is None:
                    tblHeader = OxmlElement('w:tblHeader')
                    trPr.append(tblHeader)

            c_idx = 0
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                gridSpan = tcPr.find(f'{{{ns}}}gridSpan')
                span = int(gridSpan.get(qn('w:val'), '1')) if gridSpan is not None else 1

                if c_idx < len(new_w_dxa):
                    cell_w_dxa = sum(new_w_dxa[c_idx:c_idx + span])
                    tcW = tcPr.find(f'{{{ns}}}tcW')
                    if tcW is None:
                        tcW = OxmlElement('w:tcW')
                        tcPr.append(tcW)
                    tcW.set(qn('w:w'), str(int(cell_w_dxa)))
                    tcW.set(qn('w:type'), 'dxa')
                c_idx += span

                for p in cell.paragraphs:
                    pPr = p._p.get_or_add_pPr()
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)

                    # Simpan alignment yang ada jika sudah center/right
                    # Jika di header baris, alignment center
                    if r_idx < header_rows_count:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    for r in p.runs:
                        r.font.name = font_name
                        r.font.size = Pt(table_font_sz)
                        r.bold = False

                        rPr = r._r.get_or_add_rPr()
                        for tag in ['b', 'bCs', 'color', 'highlight', 'shd', 'rStyle']:
                            el = rPr.find(f'{{{ns}}}{tag}')
                            if el is not None:
                                rPr.remove(el)
                        b = OxmlElement('w:b'); b.set(qn('w:val'), '0')
                        bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '0')
                        col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
                        rPr.extend([b, bCs, col])

    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print("Saved optimized docx to:", out)

if __name__ == '__main__':
    test_transform()
