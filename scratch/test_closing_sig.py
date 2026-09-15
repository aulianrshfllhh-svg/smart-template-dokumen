import docx
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

def clean_and_inject_signature(doc, font_name="Bookman Old Style"):
    ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    
    # Cari paragraf-paragraf di bagian akhir dokumen yang merupakan signature/pengundangan lama
    # Kata kunci: BUPATI CIREBON, SEKRETARIS DAERAH, Diundangkan di Sumber, BERITA DAERAH, IMRON, HENDRA NIRMALA
    body = doc.element.body
    
    # Hapus paragraf signature lama di ujung dokumen
    while len(doc.paragraphs) > 0:
        last_p = doc.paragraphs[-1]
        txt = last_p.text.strip().upper()
        if not txt:
            # Hapus paragraf kosong di akhir
            p_elem = last_p._p
            p_elem.getparent().remove(p_elem)
            continue
        
        is_old_sig = (
            'BUPATI CIREBON' in txt or
            'SEKRETARIS DAERAH' in txt or
            'DIUNDANGKAN DI SUMBER' in txt or
            'PADA TANGGAL' in txt or
            'BERITA DAERAH' in txt or
            'BERITA  DAERAH' in txt or
            'IMRON' in txt or
            'HENDRA NIRMALA' in txt
        )
        
        if is_old_sig:
            p_elem = last_p._p
            p_elem.getparent().remove(p_elem)
        else:
            break

    # Pastikan paragraf terakhir dari isi BAB Penutup memiliki keep_with_next = True
    # agar signature tidak pernah terpisah sendirian di halaman baru
    if len(doc.paragraphs) > 0:
        last_content_p = doc.paragraphs[-1]
        last_content_p.paragraph_format.keep_with_next = True

    # Tambahkan signature resmi BUPATI CIREBON & IMRON
    # 1. Baris "BUPATI CIREBON,"
    p_bupati = doc.add_paragraph()
    p_bupati.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bupati.paragraph_format.left_indent = Mm(101.6)
    p_bupati.paragraph_format.space_before = Pt(24)
    p_bupati.paragraph_format.space_after = Pt(72) # Ruang tanda tangan ~4 baris
    p_bupati.paragraph_format.line_spacing = 1.0
    p_bupati.paragraph_format.keep_with_next = True

    r_bupati = p_bupati.add_run("BUPATI CIREBON,")
    r_bupati.font.name = font_name
    r_bupati.font.size = Pt(12)
    r_bupati.bold = False
    rPr = r_bupati._r.get_or_add_rPr()
    for tag in ['b', 'bCs', 'color', 'highlight', 'shd']:
        el = rPr.find(f'{{{ns}}}{tag}')
        if el is not None:
            rPr.remove(el)
    b = OxmlElement('w:b'); b.set(qn('w:val'), '0')
    bCs = OxmlElement('w:bCs'); bCs.set(qn('w:val'), '0')
    col = OxmlElement('w:color'); col.set(qn('w:val'), '000000')
    rPr.extend([b, bCs, col])

    # 2. Baris "IMRON"
    p_imron = doc.add_paragraph()
    p_imron.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_imron.paragraph_format.left_indent = Mm(101.6)
    p_imron.paragraph_format.space_before = Pt(0)
    p_imron.paragraph_format.space_after = Pt(0)
    p_imron.paragraph_format.line_spacing = 1.0
    p_imron.paragraph_format.keep_with_next = False

    r_imron = p_imron.add_run("IMRON")
    r_imron.font.name = font_name
    r_imron.font.size = Pt(12)
    r_imron.bold = False
    rPr2 = r_imron._r.get_or_add_rPr()
    for tag in ['b', 'bCs', 'color', 'highlight', 'shd']:
        el = rPr2.find(f'{{{ns}}}{tag}')
        if el is not None:
            rPr2.remove(el)
    b2 = OxmlElement('w:b'); b2.set(qn('w:val'), '0')
    bCs2 = OxmlElement('w:bCs'); bCs2.set(qn('w:val'), '0')
    col2 = OxmlElement('w:color'); col2.set(qn('w:val'), '000000')
    rPr2.extend([b2, bCs2, col2])

    print("Successfully injected clean closing signature block.")
