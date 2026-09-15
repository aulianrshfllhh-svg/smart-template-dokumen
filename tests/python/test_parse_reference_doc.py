#!/usr/bin/env python3
"""
test_parse_reference_doc.py — Unit test untuk parse_reference_doc.py

Membuat sample .docx sederhana in-memory (2 BAB, 1 tabel 3 kolom)
dan memverifikasi output JSON parser.

Usage:
    python -m pytest tests/python/test_parse_reference_doc.py -v
    # atau
    python tests/python/test_parse_reference_doc.py
"""

import sys
import os
import json
import tempfile
import unittest

# Tambahkan root project ke sys.path agar bisa import scripts
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
SCRIPTS_DIR = os.path.join(PROJECT_ROOT, 'scripts')
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

try:
    import docx
    from docx.shared import Pt
    from parse_reference_doc import parse_document
except ImportError as e:
    print(f"SKIP: Dependensi tidak tersedia — {e}")
    sys.exit(0)


def create_sample_docx_with_headings(path: str):
    """
    Buat sample .docx sederhana dengan:
    - BAB I: style Heading 1 + judul "PENDAHULUAN"
      - Sub-Bab 1.1 Latar Belakang
      - Sub-Bab 1.2 Landasan Hukum
    - BAB II: teks match regex "BAB II" + judul "EVALUASI"
      - Sub-Bab 2.1 Evaluasi Pelaksanaan
    - Satu tabel 3 kolom di BAB II
    """
    doc = docx.Document()

    # --- BAB I: menggunakan Heading 1 style ---
    h1 = doc.add_paragraph('BAB I PENDAHULUAN', style='Heading 1')

    # Sub-Bab 1.1
    doc.add_paragraph('1.1 Latar Belakang')
    doc.add_paragraph(
        'Sesuai Permendagri No. 86 Tahun 2017, setiap SKPD harus menyusun Renja.'
    )

    # Sub-Bab 1.2
    doc.add_paragraph('1.2 Landasan Hukum')
    doc.add_paragraph('UU No. 25 Tahun 2004 tentang SPPN.')

    # --- BAB II: menggunakan teks plain (regex match) ---
    doc.add_paragraph('BAB II')
    doc.add_paragraph('EVALUASI RENJA SKPD')

    # Sub-Bab 2.1
    doc.add_paragraph('2.1 Evaluasi Pelaksanaan Renja Perangkat Daerah Tahun Lalu')
    doc.add_paragraph('Evaluasi mencakup program dan kegiatan tahun sebelumnya.')

    # --- Tabel 3 kolom ---
    table = doc.add_table(rows=2, cols=3)
    table.style = 'Table Grid'

    # Header baris 0
    table.cell(0, 0).text = 'Kode Rekening'
    table.cell(0, 1).text = 'Program, Kegiatan, dan Sub Kegiatan'
    table.cell(0, 2).text = 'Jumlah Dana'

    # Data baris 1
    table.cell(1, 0).text = '1.01.01'
    table.cell(1, 1).text = 'Program Perencanaan'
    table.cell(1, 2).text = '500.000.000'

    doc.save(path)
    return path


def create_sample_docx_no_bab(path: str):
    """Buat .docx tanpa heading BAB sama sekali — harus trigger error."""
    doc = docx.Document()
    doc.add_paragraph('Ini adalah dokumen biasa tanpa heading BAB.')
    doc.add_paragraph('Lorem ipsum dolor sit amet.')
    doc.save(path)
    return path


def create_sample_docx_with_table_columns(path: str, num_cols: int):
    """Buat .docx dengan tabel berkolom banyak untuk uji edge case."""
    doc = docx.Document()
    doc.add_paragraph('BAB IV', style='Heading 1')
    doc.add_paragraph('4.1 Rencana Kerja dan Pendanaan')

    table = doc.add_table(rows=1, cols=num_cols)
    table.style = 'Table Grid'
    headers = [f'Kolom {i+1}' for i in range(num_cols)]
    for col_idx, header in enumerate(headers):
        table.cell(0, col_idx).text = header

    doc.save(path)
    return path


class TestParseBabSubBab(unittest.TestCase):
    """Test deteksi BAB dan Sub-Bab."""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def test_deteksi_2_bab(self):
        """Harus mendeteksi tepat 2 BAB."""
        path = os.path.join(self.tmpdir, 'sample_2bab.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)

        self.assertIn('bab', result)
        self.assertEqual(len(result['bab']), 2, f"Diharapkan 2 BAB, dapat: {len(result['bab'])}")

    def test_nomor_bab_benar(self):
        """Nomor BAB harus 'I' dan 'II'."""
        path = os.path.join(self.tmpdir, 'sample_nomor.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)
        nomors = [b['nomor'] for b in result['bab']]

        self.assertIn('I', nomors)
        self.assertIn('II', nomors)

    def test_sub_bab_bab_pertama(self):
        """BAB I harus punya 2 sub-bab (1.1 dan 1.2)."""
        path = os.path.join(self.tmpdir, 'sample_subbab.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)
        bab_i = next((b for b in result['bab'] if b['nomor'] == 'I'), None)

        self.assertIsNotNone(bab_i, "BAB I tidak ditemukan")
        self.assertEqual(len(bab_i['sub_bab']), 2,
                         f"BAB I diharapkan 2 sub-bab, dapat: {len(bab_i['sub_bab'])}")

    def test_kode_sub_bab(self):
        """Kode sub-bab harus '1.1' dan '1.2'."""
        path = os.path.join(self.tmpdir, 'sample_kode.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)
        bab_i = next((b for b in result['bab'] if b['nomor'] == 'I'), None)
        kodes = [s['kode'] for s in bab_i['sub_bab']]

        self.assertIn('1.1', kodes)
        self.assertIn('1.2', kodes)

    def test_tipe_konten_default_rich_text(self):
        """Tipe konten default untuk sub-bab harus 'rich_text'."""
        path = os.path.join(self.tmpdir, 'sample_tipe.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)
        for bab in result['bab']:
            for sub in bab['sub_bab']:
                self.assertEqual(sub['tipe_konten'], 'rich_text',
                                 f"Sub-bab {sub['kode']} seharusnya rich_text")


class TestParseTabel(unittest.TestCase):
    """Test deteksi tabel."""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def test_deteksi_1_tabel(self):
        """Harus mendeteksi tepat 1 tabel."""
        path = os.path.join(self.tmpdir, 'sample_tabel.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)

        self.assertIn('tabel_terdeteksi', result)
        self.assertEqual(len(result['tabel_terdeteksi']), 1,
                         f"Diharapkan 1 tabel, dapat: {len(result['tabel_terdeteksi'])}")

    def test_jumlah_kolom_tabel(self):
        """Tabel harus memiliki 3 kolom."""
        path = os.path.join(self.tmpdir, 'sample_kolom.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)
        tabel = result['tabel_terdeteksi'][0]

        self.assertEqual(tabel['jumlah_kolom'], 3,
                         f"Diharapkan 3 kolom, dapat: {tabel['jumlah_kolom']}")

    def test_header_tabel(self):
        """Header tabel harus berisi 3 kolom yang benar."""
        path = os.path.join(self.tmpdir, 'sample_header.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)
        tabel = result['tabel_terdeteksi'][0]

        self.assertIn('Kode Rekening', tabel['header'])
        self.assertIn('Jumlah Dana', tabel['header'])

    def test_lokasi_bab_tabel(self):
        """Tabel yang ada di BAB II harus menunjuk ke 'BAB II'."""
        path = os.path.join(self.tmpdir, 'sample_lokasi.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)
        tabel = result['tabel_terdeteksi'][0]

        self.assertEqual(tabel['lokasi_bab'], 'BAB II',
                         f"Diharapkan 'BAB II', dapat: '{tabel['lokasi_bab']}'")

    def test_has_rowspan_colspan_false(self):
        """Tabel sederhana tanpa merge harus has_rowspan_colspan=False."""
        path = os.path.join(self.tmpdir, 'sample_no_merge.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)
        tabel = result['tabel_terdeteksi'][0]

        self.assertFalse(tabel['has_rowspan_colspan'],
                         "Tabel sederhana seharusnya has_rowspan_colspan=False")


class TestEdgeCases(unittest.TestCase):
    """Test edge case."""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def test_dokumen_tanpa_bab_raise_error(self):
        """Dokumen tanpa heading BAB harus raise ValueError."""
        path = os.path.join(self.tmpdir, 'no_bab.docx')
        create_sample_docx_no_bab(path)

        with self.assertRaises(ValueError) as ctx:
            parse_document(path)

        self.assertIn('tidak terdeteksi', str(ctx.exception).lower())

    def test_file_tidak_ada(self):
        """File yang tidak ada harus raise ValueError."""
        with self.assertRaises(ValueError):
            parse_document('/path/yang/tidak/ada.docx')

    def test_file_bukan_docx(self):
        """File bukan .docx harus raise ValueError."""
        # Buat file txt palsu
        path = os.path.join(self.tmpdir, 'bukan_docx.txt')
        with open(path, 'w') as f:
            f.write('ini bukan docx')

        with self.assertRaises(ValueError) as ctx:
            parse_document(path)

        self.assertIn('.docx', str(ctx.exception))

    def test_output_json_serializable(self):
        """Output harus JSON-serializable tanpa error."""
        path = os.path.join(self.tmpdir, 'sample_json.docx')
        create_sample_docx_with_headings(path)

        result = parse_document(path)
        # Ini tidak boleh throw exception
        json_str = json.dumps(result, ensure_ascii=False)
        self.assertIsInstance(json_str, str)
        self.assertGreater(len(json_str), 10)

    def test_tabel_banyak_kolom_masih_terdeteksi(self):
        """Tabel dengan 15 kolom harus tetap terdeteksi (flagging di PHP layer)."""
        path = os.path.join(self.tmpdir, 'sample_15cols.docx')
        create_sample_docx_with_table_columns(path, 15)

        result = parse_document(path)

        # Tabel tetap ada di output
        self.assertGreater(len(result['tabel_terdeteksi']), 0)
        tabel = result['tabel_terdeteksi'][0]
        self.assertEqual(tabel['jumlah_kolom'], 15)


if __name__ == '__main__':
    print("=" * 60)
    print("Running unit tests untuk parse_reference_doc.py")
    print("=" * 60)
    unittest.main(verbosity=2)
